from __future__ import annotations

import os
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch


NEXTGEN_SRC = Path(__file__).resolve().parents[2] / "nextgen" / "src"
if str(NEXTGEN_SRC) not in sys.path:
    sys.path.insert(0, str(NEXTGEN_SRC))

from krx_research_nextgen.models.screening import RepeatedScreenResult
from krx_research_nextgen.services.screening_export_service import ScreeningExportService


def _row(count: int, code: str, name: str, industry: str) -> RepeatedScreenResult:
    return RepeatedScreenResult(
        appearance_count=count,
        stock_code=code,
        stock_name=name,
        industry=industry,
        screen_dates=("2026-10-02", "2026-10-08"),
        average_daily_pct=1.0,
        maximum_daily_pct=2.0,
        average_weekly_pct=3.0,
        maximum_weekly_pct=4.0,
        latest_trading_value=50_000_000_000,
        latest_close=10_000,
    )


class RepeatedScreenExportTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory(
            prefix="repeated-export-test-",
            dir=Path(__file__).resolve().parents[2],
        )
        self.directory = Path(self.temporary_directory.name)
        self.exporter = ScreeningExportService()

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_keeps_existing_detailed_layout(self) -> None:
        from openpyxl import load_workbook

        target = self.directory / "detailed.xlsx"
        self.exporter.export_repeated(
            [_row(2, "000001", "가나다", "제조업")],
            str(target),
            layout="detailed",
        )

        sheet = load_workbook(target).active
        self.assertEqual(sheet.title, "기간 내 다빈도 종목")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(
            [cell.value for cell in sheet[1]][:4],
            ["등장 횟수", "종목코드", "종목명", "업종"],
        )
        self.assertEqual(
            [cell.value for cell in sheet[2]][:4],
            ["2", "000001", "가나다", "제조업"],
        )

    def test_supports_frequency_group_layout(self) -> None:
        from openpyxl import load_workbook

        target = self.directory / "frequency-groups.xlsx"
        rows = [
            _row(2, "000004", "라마바", "금융업"),
            _row(4, "000002", "다나다", "서비스업"),
            _row(3, "000003", "나나다", "운송업"),
            _row(4, "000001", "가나다", "제조업"),
        ]
        self.exporter.export_repeated(rows, str(target), layout="frequency_groups")

        sheet = load_workbook(target).active
        self.assertEqual(sheet.title, "기간 내 다빈도 종목")
        self.assertEqual(sheet.freeze_panes, "B3")
        self.assertEqual(
            set(map(str, sheet.merged_cells.ranges)),
            {"A1:A2", "B1:C1", "D1:E1", "F1:G1"},
        )
        self.assertEqual(
            [cell.value for cell in sheet[1]],
            ["횟수", "4TH", None, "3TH", None, "2TH", None],
        )
        self.assertEqual(
            [cell.value for cell in sheet[2]],
            [None, "종목명", "업종", "종목명", "업종", "종목명", "업종"],
        )
        self.assertEqual(
            [cell.value for cell in sheet[3]],
            [None, "가나다", "제조업", "나나다", "운송업", "라마바", "금융업"],
        )
        self.assertEqual([cell.value for cell in sheet[4]][:3], [None, "다나다", "서비스업"])
        self.assertEqual(sheet.column_dimensions["A"].width, 10)
        self.assertEqual(sheet.column_dimensions["B"].width, 20)
        self.assertEqual(sheet.column_dimensions["C"].width, 30)

    def test_rejects_unknown_layout(self) -> None:
        target = self.directory / "unknown.xlsx"
        with self.assertRaisesRegex(ValueError, "지원하지 않는"):
            self.exporter.export_repeated([], str(target), layout="unknown")


class RepeatedScreenExportGuiTest(unittest.TestCase):
    def test_user_can_select_each_export_layout(self) -> None:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        try:
            from PyQt6.QtWidgets import QApplication
            from krx_research_nextgen.models.screening import DateRange
            from krx_research_nextgen.windows import repeated_screen_window as module
        except ImportError as exc:
            self.skipTest(str(exc))

        class Controller:
            def __init__(self) -> None:
                self.calls = []

            def export_repeated(self, rows, path, *, layout="detailed") -> None:
                self.calls.append((rows, path, layout))

        class ImmediateTasks:
            @staticmethod
            def submit(action, *, on_result, on_error, on_finished, channel) -> None:
                try:
                    on_result(action())
                except Exception as exc:  # pragma: no cover - failure is shown by the assertion
                    on_error(str(exc))
                finally:
                    on_finished()

        application = QApplication.instance() or QApplication([])
        controller = Controller()
        window = module.RepeatedScreenWindow(
            controller,
            [],
            "주간 상승률",
            DateRange(date(2026, 6, 1), date(2026, 10, 12)),
        )
        window._tasks = ImmediateTasks()

        choices = (
            ("기존 상세 목록", "detailed", "기간내_다빈도_종목.xlsx"),
            ("횟수별 가로 배치", "frequency_groups", "기간내_다빈도_종목_횟수별.xlsx"),
        )
        for label, layout, filename in choices:
            with self.subTest(layout=layout):
                with (
                    patch.object(module.QInputDialog, "getItem", return_value=(label, True)),
                    patch.object(module.QFileDialog, "getSaveFileName", return_value=(filename, "")),
                ):
                    window._export()
                self.assertEqual(controller.calls[-1][1:], (filename, layout))

        window.close()
        application.processEvents()


if __name__ == "__main__":
    unittest.main()
